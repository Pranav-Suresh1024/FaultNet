"""Build the final editable course report from the preserved user revision and verified aggregates."""
from pathlib import Path
import json, re, sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[1]
TMP=ROOT/'.codex/final-report'
EVID=ROOT/'docs/reports/evidence'
S=json.loads((EVID/'summary.json').read_text())
L=json.loads((ROOT/'docs/reports/scholarly-references.json').read_text(encoding='utf-8'))
P=json.loads((EVID/'packet-summary.json').read_text())
provenance=json.loads((EVID/'provenance.json').read_text())
assert S['totals']['workbenchRuns']==100 and len(S['live'])==8
doc=Document(TMP/'user-source.docx' if (TMP/'user-source.docx').exists() else ROOT/'docs/reports/FaultNet-partial-report.docx')
first=next(p for p in doc.paragraphs if p.text.strip()=='TABLE OF CONTENTS')
body=doc.element.body
start=list(body).index(first._p)
for element in list(body)[start:]:
    if element.tag!=qn('w:sectPr'):body.remove(element)
for name,size in [('Normal',12),('Title',18),('Heading 1',14),('Heading 2',12),('Heading 3',12),('Caption',10)]:
    style=doc.styles[name];style.font.name='Times New Roman';style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
    fonts=style.element.get_or_add_rPr().rFonts
    for attribute in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
        fonts.attrib.pop(qn('w:'+attribute),None)
    for attribute in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+attribute),'Times New Roman')
    style.paragraph_format.space_after=Pt(6);style.paragraph_format.line_spacing=1.15
    if name.startswith('Heading'):style.font.bold=True;style.paragraph_format.keep_with_next=True;style.paragraph_format.space_before=Pt(10)
    for border in style.element.xpath('.//w:pBdr'):border.getparent().remove(border)
doc.styles['Normal'].paragraph_format.widow_control=True
doc.core_properties.author='Pranav Suresh and Soham Sameer Borikar'
doc.core_properties.last_modified_by=''
doc.core_properties.title='FaultNet Network Resilience Testing using Controlled Fault Injection'
headings=[];table_num=0;figure_num=0
def n(v,d=2):return 'Not applicable' if v is None else f'{v:,.{d}f}'
def para(text='',bold=False,style=None,align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p=doc.add_paragraph(style=style);p.alignment=align;p.add_run(text).bold=bold;return p
def page():doc.add_page_break()
def h(text,level=1,new=False):
    p=doc.add_paragraph(text,style=f'Heading {level}')
    p.paragraph_format.page_break_before=new;p.alignment=WD_ALIGN_PARAGRAPH.LEFT;headings.append(text);return p
def caption(text):
    p=para(text,style='Caption',align=WD_ALIGN_PARAGRAPH.CENTER);p.paragraph_format.keep_with_next=False;return p
def table(headers,rows,widths=None,title=None):
    global table_num
    table_num+=1
    if title:
        p=caption(f'Table {table_num}. {title}');p.paragraph_format.keep_with_next=True
    if widths is None:widths=[6.35/len(headers)]*len(headers)
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    b=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'5');e.set(qn('w:color'),'D9D9D9');b.append(e)
    t._tbl.tblPr.append(b)
    for i,values in enumerate([headers]+rows):
        row=t.rows[0] if i==0 else t.add_row();row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if i==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for j,text in enumerate(values):
            c=row.cells[j];c.width=Inches(widths[j]);c.text=str(text);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');m.append(e)
            pr.append(m)
            if i==0:s=OxmlElement('w:shd');s.set(qn('w:fill'),'EAEAEA');pr.append(s)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.05;p.alignment=WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(10);r.bold=i==0
    p=para();p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1;p.add_run().font.size=Pt(3)
    return t
def image(path,text,width=6.1):
    global figure_num
    figure_num+=1
    p=para(align=WD_ALIGN_PARAGRAPH.CENTER);p.paragraph_format.keep_with_next=True
    pic=p.add_run().add_picture(str(path),width=Inches(width))
    pic._inline.docPr.set('descr',text)
    caption(f'Figure {figure_num}. {text}')
    return pic
def link(label,url,display=None):
    p=para(label,align=WD_ALIGN_PARAGRAPH.LEFT);el=OxmlElement('w:hyperlink');el.set(qn('r:id'),doc.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=display or url;r.append(t);el.append(r);p._p.append(el)
    return p
def scenario(name):return S['scenarios'][name]

tocp=para('TABLE OF CONTENTS',True,align=WD_ALIGN_PARAGRAPH.CENTER)
tocp.paragraph_format.page_break_before=True;tocp.paragraph_format.space_after=Pt(12)
p=para();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'TOC \\o "1-3" \\h \\z \\u')
r=OxmlElement('w:r');text=OxmlElement('w:t');text.text='Table of contents';r.append(text);field.append(r);p._p.append(field)

h('1. Introduction',new=True)
para('Reliable communication requires more than a successful request under normal conditions. A service can return correct data while delivering an unacceptable response time, and a connection can end before its body is complete. Developers need controlled conditions that expose these differences and evidence that connects a configuration to its observed effect.')
para('FaultNet is a local resilience-testing system with two complementary interfaces. Its experiment workbench runs a fixed HTTP workload through baseline, faulty and recovery phases. Its live proxy cockpit applies operator-selected effects to development requests. Wireshark independently observes the real packet exchange, while separate Cisco Packet Tracer labs examine reachability after a simulated router-link failure.')
para('The central contribution is a reproducible connection between configuration, measurement and interpretation. The evaluation measures fault calibration, request completion, payload goodput, timeout tradeoffs and restoration. It also reports the cost of interception and the conditions under which the configured delivery rate is not attained. The project does not assume that a testing tool makes a connection faster; it makes failure behavior inspectable.')
h('1.1 Practical Applications',2)
para('Potential applications include checking a registration portal under delayed responses, examining a checkout interface after an HTTP 503 response, testing a download client under restricted delivery or disconnection, and exercising a frontend against controlled JSON responses. A development client must actually route its requests through the live entry for these controls to act. The implemented experiments use synthetic data rather than real customer transactions.')
para('For teaching, the workbench provides comparable phase observations and saved evidence. Packet analysis connects an HTTP result with packet timing, and the routing labs illustrate how an alternate physical path becomes useful only when a routing method selects it. These activities link application behavior, transport evidence and network-layer reachability without conflating them.')

h('2. Problem Statement')
para('Ordinary application testing often establishes that a response is correct without measuring how the same workload behaves when communication is slow, rate-limited or interrupted. Uncontrolled repetitions can change the timeout, payload or request count and produce misleading comparisons. A dashboard can also hide failures if it presents only successful timing samples.')
para('The project addresses four questions: whether the configured faults produce measurable changes; whether deadline choices affect completion; whether fault clearance restores successful responses; and whether the supporting packet and routing evidence is consistent with the claimed behavior. Measurements must preserve failed outcomes, sample counts and units as well as successful-response statistics.')

h('3. Objectives')
for text in [
'Establish a fixed 65,536-byte workload and compare its direct and proxied execution.',
'Quantify added response time and its deviation from configured delays.',
'Compare measured payload goodput with configured response-body restrictions.',
'Measure request success, failure duration and recovery evidence without replacing missing successful timings by zero.',
'Evaluate the completion-versus-waiting tradeoff while changing only the client timeout.',
'Demonstrate manual live-proxy controls, independent packet evidence and static/RIP/OSPF link-failure behavior.',
'Preserve settings and observations so that report tables and figures can be reproduced.'
]:para(text,style='List Bullet',align=WD_ALIGN_PARAGRAPH.LEFT)

h('4. Literature Review and Existing System')
for text in L['literatureParagraphs']:para(text)
table(['Study or system','Layer and faults','Automation and evidence','Limitations'],L['capabilityRows'],[1.05,1.65,1.85,1.8],title='Capability comparison with published research')

d1500=scenario('delay-1500');b16=scenario('bandwidth-16');b128=scenario('bandwidth-128')
h('5. Abstract')
para(f'FaultNet is a local HTTP resilience-testing system that combines automated experiments with manual development-proxy controls. The workbench coordinates baseline, faulty and recovery phases against a fixed 64 KiB response, validates observations and saves JSON evidence. The evaluation comprises {S["totals"]["workbenchRuns"]} workbench runs with {S["totals"]["workbenchAttempts"]} phase observations, 100 direct/proxy calibration requests and 80 live-proxy samples. Ten repeated runs are used per workbench configuration. The 1,500 ms delay condition adds a median {n(d1500["addedMedianMs"])} ms relative to matching baselines. At 16 KiB/s, observed payload goodput is {n(b16["goodputKiBps"])} KiB/s, while the 128 KiB/s condition attains {n(b128["rateAttainmentPct"])}% of its configured rate. Increasing timeout from 500 to 2,000 ms under a fixed 1,000 ms delay changes faulty completion from {n(scenario("timeout-500")["availabilityPct"],0)}% to {n(scenario("timeout-2000")["availabilityPct"],0)}%. Packet timings provide independent traffic evidence. Reviewed routing cases show that RIP and OSPF use the alternate path after the tested direct-link failure, whereas the static control lacks a backup route. The findings support controlled local fault testing and expose platform-dependent pacing limitations.')

h('6. Workflow of the Proposed Project work and System Architecture')
para('A single Node.js process owns four loopback listeners. Port 3000 serves both interfaces and their protected control APIs. The fixed demo on port 3001 returns HTTP 200 and exactly 65,536 synthetic F bytes for GET /payload. The experiment proxy on port 3002 applies the workbench fault. The separate live proxy on port 8080 forwards development traffic under its applied configuration.')
table(['Component','Port','Function'],[('Control server','3000','Dashboard assets, settings, readiness, records, exports and live control.'),('Demo server','3001','Fixed response contract for repeatable experiments.'),('Experiment proxy','3002','One selected delay, rate restriction or interruption per run.'),('Live proxy','8080','Independent configurable HTTP/HTTPS development forwarding.')],[1.4,.6,4.35],title='Listener responsibilities')
para('The measured workbench path is Node probe → experiment proxy → demo server; the body returns through the proxy to the probe. The runner assigns phases, clears faults and saves observations. The browser retrieves measurements through the control API and renders charts; it neither generates the measured duration nor directly reads result files. Control traffic is kept separate from the faulted data path.')
image(EVID/'figures/architecture.png','Control and data paths in FaultNet. The workbench and live proxy have independent fault configurations; replies return through the same proxy path.')
para('The workbench validates a request, saves its initial record, runs baseline probes, activates the selected fault, runs faulty probes, clears the fault, checks recovery and saves the terminal record. One experiment is active at a time. Stop cancels active work and clears held operations. If a process ends unexpectedly, a previously active persisted record is marked interrupted on restart rather than silently completed.')
h('6.1 Dashboard Roles and Separation',2)
table(['Aspect','Experiment workbench','Live proxy cockpit'],[
('Question','How does this fault alter a controlled workload?','How do individual development requests behave under applied effects?'),
('Target','Fixed 64 KiB demo','Validated development origins; demo by default'),
('Sequence','Automatic baseline/fault/recovery','Operator applies settings, sends traffic and resets'),
('Persistence','Validated local JSON and JSON/CSV exports','Latest 50 finished requests in process memory'),
('Measurements','Phase counts, distributions and comparisons','Per-request status, duration, bytes and applied effects')
],[1.0,2.65,2.7],title='Why the project uses two dashboards')
para('Separate configurations protect reproducible measurements from exploratory edits. The workbench constrains the workload and comparison protocol; the cockpit provides broader manual combinations. Changing one dashboard does not change the other proxy\'s fault configuration. Neither interface automatically repairs a failing application.')
h('6.2 Workbench Functions and Fault Mechanisms',2)
para('Six editable presets cover a healthy baseline, mild/severe delay, constrained bandwidth, timeout stress and complete HTTP outage. A preset fills settings without starting traffic. Start begins the experiment; Stop requests cancellation; saved-run reuse copies a configuration without modifying its evidence. Auto or explicit phase selection controls the displayed analysis. Protocol-compatible saved runs can be compared, and records can be exported as JSON or CSV.')
table(['Mechanism','Implementation action','Measured consequence'],[('Delay','Await a timer before forwarding','Additional complete-response time; possible client timeout.'),('Bandwidth','Schedule 1 KiB body chunks at rate-based deadlines','Longer transfer and lower payload goodput.'),('Interruption','Destroy the faulty client connection before forwarding','No complete successful HTTP payload; subsequent recovery checked after clearance.')],[1.0,2.65,2.7],title='The three workbench fault mechanisms')
para('The timer wait is asynchronous and does not busy-wait. Bandwidth scheduling honors downstream backpressure and buffers only the known payload. These operations alter actual proxied HTTP exchanges, not the laptop\'s physical bandwidth or an IP packet-loss percentage. Timeout stress is an existing delay mechanism paired with a shorter client deadline.')
para('Each probe succeeds only after HTTP 200, all 65,536 bytes and the expected payload contents arrive before timeout. Mean, median and p95 use successful samples. Eligible requests exclude user cancellation, and unavailable successful timing remains null. Charts present all three phase medians, attempt durations, successful-response distributions and successful/failed/cancelled counts.')
h('6.3 Live Proxy Cockpit Functions',2)
para('Unified routing forwards to one origin; split routing selects frontend or backend origins using path prefixes. The client must use the port-8080 entry for effects to apply. Presets/sliders change a draft, Apply changes activates it for new requests, and Reload applied retrieves current settings. Reset removes effects and rules, discards drafts and cancels live work while retaining the applied targets and routes.')
table(['Control','Purpose'],[('Request/response delay and jitter','Vary forwarding and response waits; jitter applies in both directions and clamps negative waits to zero.'),('Upload/download limits','Pace body delivery in KiB/s; zero means unlimited.'),('Timeout / stall / disconnect','Delay headers, hold a body, or interrupt body delivery.'),('Status injection','Return a selected 400–599 status for a path prefix at a configured probability.'),('Header controls','Remove CORS/cache headers or mismatch Content-Type.'),('JSON mocks','Return local JSON without contacting upstream; selected status rules have precedence.')],[1.8,4.55],title='Live fault and response controls')
para('Send sample performs a GET with a five-second client ceiling. The monitor provides status, duration, bytes, outcome, effect labels, filtering and aggregate counts for the latest 50 finished requests. Clear monitor affects that window only. Live configuration and monitoring survive a page refresh but reset with the process. Named mobile profiles are illustrative settings, not carrier measurements.')

h('7. Contribution of Individual Team Member')
table(['Member','Registration number','Project responsibilities'],[
('Pranav Suresh','24BCE2086','Coordinated workbench/cockpit requirements, presets and interface refinements; prepared HTTP/Wireshark demonstrations; coordinated repeated evaluation, interpretation and the editable report/evidence package.'),
('Soham Sameer Borikar','24BAI0142','Prepared and uploaded static, RIPv2 and OSPF .pkt labs on the separate lab branch; configured the router triangle and alternate paths; prepared the routing/link-failure demonstration.')
],[1.55,1.05,3.75],title='Division of project responsibilities')
para('The application and routing components form one demonstration with distinct evidence paths. Both members must be able to explain the fault settings, actual observations, routing changes and limitations of the conclusions. The source repository preserves application context and the separately reviewed lab history.')

h('8. Modules Explanation')
table(['Important module','Responsibility'],[
('src/app.js and src/main.js','Dependency wiring, listener startup and orderly shutdown.'),
('src/control/server.js','API routing for experiments, records, export and live controls.'),
('src/control/validation.js and security.js','Input bounds, allowed shapes, Host/Origin checks and capability protection.'),
('src/experiments/runner.js','Phase coordination, cancellation, deadline, recovery and saves.'),
('src/experiments/probe.js','HTTP request, elapsed timer, payload/status verification and outcomes.'),
('src/experiments/metrics.js','Phase counts, mean/median/p95 and payload goodput.'),
('src/proxy/server.js','Fixed-target delay, body pacing and interruption.'),
('src/live/server.js and config.js','Independent forwarding, validated targets and live effects.'),
('src/storage/records.js and export.js','Record validation, safe persistence and JSON/CSV export.'),
('public/app.js, results.js and charts.js','Workbench events, derived views, comparisons and SVG visualizations.'),
('public/live.js and live-model.js','Applied/draft controls, presets and live traffic presentation.')
],[2.85,3.5],title='Important code modules')

h('9. Project Implementation')
h('9.1 Experimental Setup with Hardware and Software Requirements',2)
env=S['environment'];para(f'The application evaluation used {env["os"]} build {env["release"]}, a {env["cpu"]} processor and {n(env["memoryGiB"],1)} GiB of visible memory. The measured application runtime was {S["runtime"]}. The application uses Node built-in HTTP, timing, filesystem and test facilities with plain HTML/CSS/JavaScript and no external runtime packages. Analysis tooling is separate from application dependencies.')
para(L['reproducibilityParagraph'])
para('All measured traffic used the numeric loopback address. HTTP probes disabled connection pooling, verified the fixed payload, and measured the entire exchange with performance.now(). UTC starts identify records for packet matching. Background desktop applications and normal operating-system scheduling remained present; the results describe this local environment rather than an isolated hardware benchmark.')
link('Project repository: ','https://github.com/Pranav-Suresh1024/FaultNet')
para('Application revision: '+provenance['applicationCommit']+'. Evaluation scripts and sanitized numerical inputs: '+provenance['evidenceCommit']+'. Appendix A separates document regeneration, numerical recalculation and fresh collection.')
link('Permanent evidence snapshot: ',provenance['evidenceUrl'],'Evaluation scripts, numerical inputs and aggregate evidence at the pinned revision')
table(['Protocol choice','Value or procedure'],[
('Repeated blocks','10 runs per configuration; cyclic configuration order across blocks.'),('Workload','65,536 F bytes; five baseline and five faulty probes per run.'),('Interval','250 ms between probes within a phase.'),('Default deadline','5,000 ms; timeout study changes only this to 500 or 2,000 ms.'),('Recovery','Three consecutive successful responses within 15,000 ms after clearance.'),('Overhead test','50 direct and 50 fault-free proxied requests, grouped in ten paired blocks; three warm-up pairs excluded.'),('Live study','Eight configurations, ten sequential samples each; five-second sample ceiling.'),('Packet capture','Npcap loopback; TCP 3001, 3002 and 8080; dashboard control excluded.')
],[1.45,4.9],title='Measurement protocol and reproducibility controls')
para('The comparison budget is 2 × count × (timeout + interval) + 15,000 ms. Five probes, a 5,000 ms timeout and 250 ms interval give a conservative 67,500 ms bound, below 120,000 ms. This is a safety bound, not a duration prediction. The timeout pair retains the same payload, delay, count and interval; its different deadlines intentionally make it incompatible with the dashboard\'s automatic cross-run guard.')
h('9.2 Test Cases',2)
table(['Condition','Controlled variable','Expected observation'],[
('Delay calibration','0, 150, 500, 1,500 ms','Added timing follows the injected wait with measurable overhead.'),('Bandwidth calibration','16, 32, 128 KiB/s','Ideal 64 KiB delivery is 4, 2 and 0.5 seconds; actual rate is measured.'),('Connection interruption','Close each faulty connection','Faulty requests fail to complete, followed by successful clearance/recovery.'),('Timeout tradeoff','1,000 ms delay; 500 vs 2,000 ms timeout','Completion may improve when the deadline exceeds the delay.'),('Live behavior','Normal, delay, jitter, rate, 503, disconnect, stall, mock','Each request exhibits the applied effect with honest status/bytes/outcome.'),('Routing link failure','R1 G0/0 disabled','Static control lacks backup; dynamic labs can select R3.')
],[1.35,2.15,2.85],title='Experimental cases and hypotheses')
para('A completed experiment is not automatically a successful faulty workload: the interruption and short-timeout conditions can complete the experiment while recording all faulty requests as failures. No cancelled request was reclassified as a successful timing sample. Run-level record validation checks structure, summaries, ordering and recovery chronology before analysis.')
para('Within each run, added delay is faulty median minus baseline median. The calibration table reports the median of ten paired run-level differences. The exploratory 95% percentile bootstrap interval uses 3,000 resamples and seed 308, resampling whole run-level differences rather than treating the five faulty requests within each run as independent repeats. There are ten runs per condition and 50 faulty requests in total. The sorted bootstrap medians use positions 75 and 2,925 (one-based). These local intervals and pooled nearest-rank p95 are descriptive, not population guarantees.')
para('Request success rate (%) equals 100 × successful attempts / eligible attempts; user cancellations are excluded from eligibility. For a successful fixed response, payload goodput (KiB/s) equals 64 / (elapsed milliseconds / 1,000); the table reports the arithmetic mean of these per-request rates. The mean of reciprocal durations need not equal 64 KiB divided by the median duration. The nearest-rank p95 uses position ceiling(0.95 × N) in sorted successful timings. These definitions preserve the distinction between slow success, failed completion and missing successful latency.')
para('The configured automated test command passed all 61 tests with zero failures or skipped tests. Coverage includes fault behavior, cancellation and cleanup, recovery, persistence/export validation, control security, live proxy behavior and frontend calculations. This verifies the tested paths; it does not establish correctness under every workload or replace the separate packet and simulator evidence.')

h('9.3 Results of Execution and Screenshots',2)
para(f'The completed dataset contains {S["totals"]["workbenchRuns"]} workbench runs and {S["totals"]["workbenchAttempts"]} phase observations, plus 100 overhead requests and 80 live samples. Each workbench condition has ten runs and 50 faulty attempts. The following tables report measurements, while the comparison plots explain their relationship to configuration. The workbench interface figure is a representative reviewed screenshot; its displayed example values are separate from the repeated-study aggregates.')
quality=S['recordQuality']
para(f'Record quality is reported separately from request success. There were {quality["accepted"]+quality["rejected"]} attempted workbench runs: {quality["accepted"]} supplied validated terminal records and {quality["rejected"]} failed record validation. In the rejected bandwidth trial, UTC recovery duration and elapsed-clock duration differed by {n(quality["rejections"][0]["utcVsElapsedDifferenceMs"])} ms, above the validator\'s 2 ms tolerance. The retained rejected record is excluded from timing aggregation and the condition is repeated to preserve the ten-run design. Thus 100 accepted runs do not imply that every attempted record was saved successfully.')
image(EVID/'screenshots/workbench.png','Representative workbench interface showing settings, phase comparisons, successful timing and recovery evidence. Displayed example values are not the repeated-study aggregate.')
para('Repeated-run summaries and figures are calculated from measured observations; the representative interface image illustrates the controls and their presentation.')
names=['delay-0','delay-150','delay-500','delay-1500','bandwidth-16','bandwidth-32','bandwidth-128','interruption','timeout-500','timeout-2000']
table(['Condition','Success','Successful median ms','Successful p95 ms','Mean per-request goodput KiB/s'],[(name,f'{scenario(name)["successes"]}/50',n(scenario(name)['medianMs']),n(scenario(name)['p95Ms']),n(scenario(name)['goodputKiBps'])) for name in names],[1.35,.65,1.15,1.15,2.05],title='Faulty-phase results across repeated workbench trials')
para('A missing successful median or goodput is not zero: interruption and deadline failures have no successful full-response sample. Their attempted durations are analysed separately. Bandwidth and delay both reduce whole-request goodput, but their mechanisms differ: one paces the body and the other waits before forwarding.')

h('9.4 Performance Analysis and Comparison of Methods',2)
h('9.4.1 Delay Calibration and Distribution',3)
table(['Delay ms','Paired added ms','Error %','95% percentile bootstrap interval ms'],[(str(scenario(name)['fault']['delayMs']),n(scenario(name)['addedMedianMs']),n(scenario(name).get('calibrationErrorPct')),f'{n(scenario(name)["addedBootstrap95Ms"][0])} to {n(scenario(name)["addedBootstrap95Ms"][1])}') for name in ['delay-0','delay-150','delay-500','delay-1500']],[.85,1.4,.85,3.25],title='Configured delay against paired added response time')
image(EVID/'figures/delay-calibration.png','Configured delay versus paired measured increase. Error bars are exploratory bootstrap intervals across ten runs.')
para(f'The 1,500 ms condition produced an added median of {n(d1500["addedMedianMs"])} ms, an error of {n(d1500["calibrationErrorPct"])}%. The zero-delay condition adds {n(scenario("delay-0")["addedMedianMs"])} ms relative to its in-run no-fault baseline. Percentage error is undefined at a zero setting. This differs from the separate {n(S["overhead"]["blockAddedMedianMs"])} ms direct-versus-fault-free-proxy estimate: both phases of a zero-delay run already traverse the proxy, and only the faulty phase dispatches the delay timer. The studies also use different sequences and background conditions. Timer dispatch and desktop scheduling are plausible contributors, but this dataset does not isolate the cause of the gap. A controlled follow-up should interleave fault-free proxy and zero-delay operations, instrument timer waits, and repeat under controlled load; 14.61 ms must not be subtracted as a universal correction.')
image(EVID/'figures/latency-distribution.png','Empirical distributions of successful faulty response times; 50 observations per configured delay. Curves show separation and within-condition variation.')
para('Distribution separation shows that the configured wait changes actual response behavior, rather than only changing a summary label. A percentile based on 50 successful samples remains descriptive. It should be read alongside the number of successes, failure outcomes and the repeated-run spread.')

h('9.4.2 Bandwidth Shaping and Rate Attainment',3)
table(['Limit KiB/s','Ideal ms','Median ms','Goodput KiB/s','Attainment %'],[(str(scenario(name)['fault']['rateKiBps']),n(64000/scenario(name)['fault']['rateKiBps']),n(scenario(name)['medianMs']),n(scenario(name)['goodputKiBps']),n(scenario(name)['rateAttainmentPct'])) for name in ['bandwidth-16','bandwidth-32','bandwidth-128']],[1.0,1.0,1.15,1.65,1.55],title='Configured response rate compared with measured payload goodput')
image(EVID/'figures/bandwidth-calibration.png','Measured whole-request goodput versus configured restriction for the same 64 KiB payload.')
para(f'The 16 KiB/s condition attained {n(b16["rateAttainmentPct"])}% of its configured rate. The 128 KiB/s condition attained {n(b128["rateAttainmentPct"])}%, so the implementation should be interpreted as a delivery cap rather than a guarantee of exact target throughput. The 1 KiB schedule has nominal chunk intervals of 62.5, 31.25 and 7.8125 ms at 16, 32 and 128 KiB/s respectively. At short waits, timer dispatch and local scheduling can consume a larger fraction of the interval. This is a plausible explanation consistent with the mechanism; timer causes were not independently isolated.')
para('This comparison is useful even where calibration is weaker: it identifies the operating range where the simple pacing approach follows its configuration closely. The report retains the high-rate result instead of silently discarding it or widening an accuracy claim.')

h('9.4.3 Timeout Tradeoff and Failure Outcomes',3)
image(EVID/'figures/timeout-tradeoff.png','Fixed 1,000 ms injected delay with two client deadlines. The duration panel distinguishes failed waiting time from successful response latency.')
short=scenario('timeout-500');long=scenario('timeout-2000')
para(f'At a 500 ms timeout, {short["successes"]} of 50 faulty requests succeeded and the median failed-attempt duration was {n(short["failureMedianMs"])} ms. At 2,000 ms, {long["successes"]} of 50 succeeded with a successful median of {n(long["medianMs"])} ms. The observed completion difference is {n(long["availabilityPct"]-short["availabilityPct"],0)} percentage points. The delay itself was unchanged. Longer waiting improves tolerance in this case, but also postpones failure detection when a service will not recover.')
para('The interruption condition produced '+str(scenario('interruption')['successes'])+' successful faulty responses out of 50. This is an application-level completion result; it is not an estimate of random IP packet loss, nor does it imply that every connection necessarily used a TCP reset flag.')

h('9.4.4 Interception Cost and Recovery',3)
o=S['overhead'];para(f'Across 50 direct and 50 fault-free proxied requests, pooled medians were {n(o["directMedianMs"])} and {n(o["proxyMedianMs"])} ms. The median paired block difference was {n(o["blockAddedMedianMs"])} ms, with an exploratory bootstrap interval from {n(o["blockAddedBootstrap95Ms"][0])} to {n(o["blockAddedBootstrap95Ms"][1])} ms. This quantifies instrumentation cost on the fixed local workload. It is not a web-scale throughput measurement.')
image(EVID/'figures/recovery-comparison.png','Baseline and recovered successful response times for selected conditions; Table 13 covers all ten. Recovery phase medians and confirmation duration are distinct metrics.')
table(['Condition','Confirmed runs','Confirmation median ms','Recovery minus baseline ms'],[(name,f'{scenario(name)["recoveryConfirmed"]}/10',n(scenario(name)['recoveryMedianMs']),n(scenario(name)['recoveryResidualMs'])) for name in names],[1.55,1.1,1.7,2.0],title='Restoration evidence for all ten conditions')
para('With three required successes and 250 ms between recovery attempts, two waits alone contribute approximately 500 ms. Confirmation therefore includes observation policy and scheduling, not just fault removal. The recovery residual compares successful response behavior after clearance with that run\'s baseline and is useful alongside the confirmation flag.')

h('9.4.5 Live Proxy Evaluation',3)
image(EVID/'screenshots/live-cockpit.png','Live proxy cockpit showing applied settings and its recent-traffic monitor. This interface snapshot is separate from the controlled evaluation tables.')
para('The cockpit snapshot shows 50 finished requests, 49 marked altered, 30 error/cancel outcomes and a 29.9 ms displayed median. These values summarize the latest monitor window, including earlier faulted requests; zero delay/rate settings do not erase that history. Reset removes active effects, while Clear monitor clears the visible history. The rolling monitor values are distinct from the eight-condition live study and the workbench\'s successful-response statistics. A display of error/cancel outcomes is not an IP packet-loss measurement.')
monitor=image(EVID/'screenshots/live-cockpit.png','Enlarged monitor detail from the same cockpit screenshot. The latest 50 finished requests remain visible after faults are reset; counts are not the repeated-study totals.')
# Preserve the original image bytes; Word's crop enlarges the monitor for print.
monitor.height=Inches(6.1*740/687)
crop=OxmlElement('a:srcRect');crop.set('l','63400');crop.set('t','15621')
monitor._inline.xpath('.//pic:blipFill')[0].insert(1,crop)
image(EVID/'figures/live-comparison.png','Individual live attempt durations under normal, delayed, jittered and rate-restricted configurations; ten observations per condition, with black median markers.')
table(['Live condition','Completed','Completed with HTTP 200','Median attempt duration ms','Median bytes'],[(name,f'{s["completed"]}/10',f'{s["http200"]}/10',n(s['medianMs']),n(s['medianBytes'],0)) for name,s in S['live'].items()],[1.35,.8,1.4,1.6,1.2],title='Manual live-proxy sample outcomes')
normal=S['live']['normal'];jitter=S['live']['jitter-100'];fixed=S['live']['delay-500']
para(f'The fixed-delay condition had a sample standard deviation of {n(fixed["sdMs"])} ms, compared with {n(jitter["sdMs"])} ms under configured jitter. Jitter applies to both request and response waits, including the request direction with base zero, where negative variation is clamped. Consequently the jitter condition is not assumed to have exactly the same mean as fixed response delay alone.')
para(f'The 100% HTTP 503 rule returned that status in {S["live"]["http-503"]["http503"]}/10 samples. Those samples can be completed HTTP exchanges while still reporting an application error. Stalled-body samples can receive headers before the five-second sample ceiling expires, so a recorded status 200 alone is insufficient to establish success. Body bytes and completion outcome must also be inspected. The final reset check returned HTTP {S["restored"]["status"]} with {S["restored"]["bytes"]:,} bytes.')

h('9.4.6 Packet Evidence and Timing Correlation',3)
para(f'The filtered loopback capture contains {P["packetCount"]:,} packets; {P["controlPackets"]} packets match the dashboard control port. Complete HTTP reconstruction on port 3002 enables stream-specific timing checks. Probe starts are matched to GET timestamps and stream endpoints, with first response data and completed reconstruction considered separately. The capture began after the overhead warm-up sequence, so it is used only for matching covered experiment requests.')
para('Capture-drop statistics are unknown: the saved capture has no interface-statistics block. Zero control-port packets describes the filter, not zero capture loss. TShark/Npcap observations establish the matched traffic intervals without asserting that every generated packet was recorded.')
table(['Condition','Matched streams','Median HTTP time ms','Median probe time ms'],[(row['scenario'],str(row['matched']),n(row['packetMedianMs']),n(row['probeMedianMs'])) for row in P['matches']],[1.5,1.1,1.9,1.85],title='Packet-to-probe correlation for representative covered runs')
image(EVID/'figures/packet-evidence.png','Measured request and response timing extracted with TShark from the loopback capture. This is an analytical packet view, not a simulated Wireshark screenshot.')
para(L['packetDiscussion'])

h('9.4.7 Routing Comparison in Packet Tracer',3)
image(EVID/'screenshots/packet-tracer-topology.png','Reviewed Packet Tracer topology with three routers, three switches and six PCs. The router triangle offers an alternate R1–R3–R2 path.')
para('The reviewed lab workload sends PC1 (192.168.0.2) to PC3 (192.168.1.2). R1 GigabitEthernet0/0 is disabled to remove the direct R1–R2 path. Stable ping batches and the R1 route table are inspected after convergence, and the interface is restored afterward. These simulator observations are a separately reviewed evidence set, not repetitions of the HTTP experiment dataset.')
para('The independent review used Cisco Packet Tracer 9.0.1 and lab revision 9f72386cf4a1ae30ccd50ba7c0a1a403b3178229. The topology image establishes layout; the following route excerpts and settled ping summaries support reachability and alternate next-hop selection.')
table(['Method','Settled baseline','Direct link disabled','Alternate-route evidence'],[('Static without backup','4/4 replies','0/4 replies','No configured alternate route for the destination LAN.'),('RIPv2','4/4 replies','4/4 replies','R 192.168.1.0/24 via 10.0.0.9, G0/1; [120/2].'),('OSPF','4/4 replies','4/4 replies','O 192.168.1.0/24 via 10.0.0.9, G0/1; [110/3].')],[1.35,1.0,1.1,2.9],title='Settled routing observations under the same direct-link failure')
image(EVID/'figures/routing-comparison.png','Descriptive reachability comparison of reviewed four-ping batches. Bars are observations of this topology, not reliability estimates for each protocol.')
for line in ['R1# show ip route  |  direct R1-R2 link disabled',
             'RIP: R 192.168.1.0/24 [120/2] via 10.0.0.9, GigabitEthernet0/1',
             'OSPF: O 192.168.1.0/24 [110/3] via 10.0.0.9, GigabitEthernet0/1',
             'PC1> ping 192.168.1.2  |  settled fault batches:',
             'Static: 0/4 replies; RIPv2: 4/4 replies; OSPF: 4/4 replies.']:
    p=para(line,align=WD_ALIGN_PARAGRAPH.LEFT);p.paragraph_format.space_after=Pt(2)
    for run in p.runs:run.font.name='Consolas';run.font.size=Pt(9)
para('The route lines are reviewed excerpts; the final ping line is a compact summary of separate batches, not one combined terminal session. Both dynamic routes use R3 (10.0.0.9) through G0/1.')
para('Dynamic routing preserved reachability in the settled fault state of this topology, while the deliberately unprotected static configuration did not. Static routing can be configured with backup routes; the result is not a universal failure of static routing. Warm-up and restoration batches included transient losses. Fast Forward Time advanced the simulator, so the records do not support wall-clock convergence comparisons or a claim that OSPF was faster than RIP. Route selection, next hop and reachability establish the demonstrated benefit.')

h('9.5 Critical Evaluation and Optimization',2)
para('The strongest evidence comes from controlled changes paired with matching baselines. Delay calibration quantifies whether the fault behaves as configured. Bandwidth attainment shows both the useful lower-rate range and the limits of fine-grained local pacing. The timeout study demonstrates a real completion-versus-waiting tradeoff without pretending that a higher deadline optimizes the network.')
para('The metrics also prevent misleading conclusions. Success counts accompany latency; unavailable successful timings are not zeros. Recovery confirmation includes the polling policy. Direct comparison measures interception overhead rather than assuming it is negligible. Routing metrics are interpreted within their protocol and are not numerically ranked across RIP and OSPF.')
para('The current 1 KiB scheduling method could be evaluated against adaptive chunk sizes or batching at higher rates. A sound optimization study would hold the payload and target rate constant, compare attained goodput and CPU use, and retain cancellation/backpressure behavior. That is an engineering implication of the measured pacing limitation, not an implemented improvement claimed by this report.')

h('10. Results and Discussion')
h('10.1 Findings and Inferences',2)
para(f'First, added response time follows the configured delay, with the 1,500 ms condition producing {n(d1500["addedMedianMs"])} ms of paired added median. Second, payload pacing is close to its lower configured restrictions but attains only {n(b128["rateAttainmentPct"])}% at 128 KiB/s. Reporting both outcomes is more informative than assuming every rate is exact.')
para(f'Third, the timeout intervention changes faulty completion by {n(long["availabilityPct"]-short["availabilityPct"],0)} percentage points while retaining the same 1,000 ms configured wait. The longer deadline tolerates that response; it does not remove its delay. Fourth, post-clearance observations supply repeated restoration evidence and distinguish service response timing from the duration of the recovery criterion.')
para('Live results show why status, body size and completion must be considered together. A mock can complete with a small local JSON body; an injected 503 can complete with an error status; a stalled body can time out after headers. The cockpit therefore addresses manual application behavior that cannot be reduced to the workbench\'s fixed payload success criterion.')
para('Packet evidence is consistent with real HTTP traffic under the tested conditions. The separately reviewed routed topology demonstrates alternate-path selection under a specific failure. The application and simulator answer related resilience questions at different layers; neither evidence set is used to invent measurements for the other.')
h('10.2 Innovation Novelty and Evidence',2)
para(L['noveltyParagraph'])
para('Useful improvement claims are attached to actual comparisons. The longer deadline improves completion for the chosen delayed workload, and dynamic route selection improves settled reachability over the no-backup control in the reviewed topology. Neither claim implies general superiority over unrelated systems. Calibration and interception results identify costs as well as benefits.')

h('11. Conclusion')
para(f'FaultNet implements controlled HTTP fault injection with saved experiments and a separate live development cockpit. A repeated evaluation of {S["totals"]["workbenchRuns"]} workbench runs supports measurable delay, pacing, failure and restoration behavior. The timeout comparison improves observed completion from {n(short["availabilityPct"],0)}% to {n(long["availabilityPct"],0)}% by permitting longer waiting, while the high-rate pacing result exposes a concrete limitation. Packet analysis and reviewed routing observations provide complementary evidence. The system is useful for local development testing and networking education because configurations, measurements and explanations can be examined together.')

h('12. Limitations and Future Enhancement')
para('The study uses one shared Windows laptop, sequential requests and a fixed synthetic payload. Operating-system scheduling, background applications and capture activity can affect timing. Fifty faulty attempts per condition support descriptive comparisons, not universal request-success or tail-latency guarantees. The bootstrap intervals describe variation among ten local blocks and do not remove environmental dependence.')
para('One attempted run failed the record validator at the UTC-versus-elapsed recovery tolerance boundary. This is a persistence/validation limitation, separate from an injected connection failure. The rejected record remains in the local audit data and its count is included in record quality; only structurally validated records contribute to the comparison tables.')
para('Request elapsed time is not TCP RTT, payload goodput is not wire bandwidth, and HTTP interruption is not random IP loss. Recovery confirmation depends on three-success sampling and its waits. The routing cases use small settled ping batches and simulated-time advancement; they establish path selection without a convergence-speed ranking. Representative interface screenshots illustrate the application rather than replacing the current measurement tables.')
para('Further evaluation can vary payload sizes, use additional machines and controlled background load, compare pacing strategies at short chunk intervals, and measure routing convergence on a common time base. More frequent crash-durable saves and separate processes could be assessed alongside their storage and coordination costs. These extensions follow directly from the observed limits and do not alter the conclusions for the measured configuration.')

h('References',new=True)
for ref in L['papers']:
    p=para(f'[{ref["number"]}] {ref["authors"]}, “{ref["title"]},” {ref["publication"]}.',align=WD_ALIGN_PARAGRAPH.LEFT);p.paragraph_format.keep_with_next=True
    link('DOI: ','https://doi.org/'+ref['doi'],ref['doi']).paragraph_format.keep_with_next=True
    link('Free full text: ',ref['fullTextUrl'],ref['fullTextVersion'])
    if ref.get('additionalFullTextUrl'):link('Additional author copy: ',ref['additionalFullTextUrl'],'Stanford-hosted full conference paper')

h('Appendix A Reproduction Procedure',new=True)
para('The pinned evidence revision in section 9.1 contains evaluation/analysis scripts and the numerical inputs used for the published results. It preserves the original measurements; recalculation is not a new experiment. The later report revision contains the document builder and the revised prose/layout.')
para('1. Regenerate the document: use the report revision identified in FINAL-REPORT-HANDOFF.md, install Python packages python-docx and pypdf, run python scripts/build-final-report.py, then use Microsoft Word to refresh the contents and export the PDF (scripts/export-report.ps1). This reads committed aggregates and figures; it does not recalculate results. Close or save any open report before rebuilding. The preserved partial report supplies the course cover and declaration.')
para('2. Recalculate numerical results: install matplotlib and numpy, then run python scripts/analyze-report.py from the repository. The default input is docs/reports/evidence/analysis-inputs.json. It contains the original numerical phase durations, success flags, recovery durations, overhead blocks and live outcomes, without session capabilities, HTTP bodies or wall-clock timestamps. It reproduces the workbench/live/overhead statistics, bootstrap intervals and their figures. The rejection count and clock discrepancy are retained audit facts; this reduced dataset cannot revalidate the original record chronology.')
para('3. Packet and routing evidence: the committed packet summary allows inspection of the 15 matched timings and frame references, but independently decoding streams or checking capture completeness requires the original local capture. Route excerpts and ping summaries document the bounded simulator review; reproducing that execution requires Packet Tracer 9.0.1 and the separately pinned lab files. Sanitized inputs do not replace packet capture or simulator execution.')
para('4. Collect replacement measurements: use Node 22.23.1 and npm start on an idle local setup, then node scripts/evaluate-report.mjs. The collector checks application files against the pinned application baseline. Preserve any rejected record and resume only after review. Run python scripts/prepare-report-inputs.py to derive a sanitized numerical dataset, then analyze it. Fresh timings will differ and must be labelled as a replacement evaluation; never silently combine new samples with this study. Follow the evidence guide for optional loopback capture and packet correlation.')
confirmation=json.loads((EVID/'followup-verification.json').read_text())
assert confirmation['tests']['passed']==61 and confirmation['workbench']['acceptedRuns']==10 and confirmation['live']['attempts']==80
para('Separate confirmation check. A fresh execution of the unchanged application passed all 61 automated tests with zero failures, skips or cancellations. One additional run of each of the ten workbench conditions produced 130 phase observations (50 baseline, 50 faulty and 30 recovery); all ten records saved successfully and confirmed recovery, with no validation rejection. Interruption and the 500 ms deadline produced 0/5 faulty successes; the remaining conditions produced 5/5. Eight live configurations were exercised with ten samples each, and the final reset returned HTTP 200 with 65,536 bytes. This bounded confirmation is recorded in evidence/followup-verification.json and is separate from the original 100-run study: it does not replace its statistics, intervals, packet capture or routing observations.')

# Retain native page fields; Word rebuilds the automatic index before export.
for p in doc.paragraphs:
    if p.style.name.startswith('Heading'):p.paragraph_format.widow_control=True
settings=doc.settings.element
for e in settings.xpath('.//w:updateFields'):settings.remove(e)
e=OxmlElement('w:updateFields');e.set(qn('w:val'),'true');settings.append(e)
assert len([x for x in headings if re.match(r'^\d+\. ',x)])==12
assert not any('yet to be implemented' in p.text.lower() or 'out of scope' in p.text.lower() for p in doc.paragraphs)
path=Path(sys.argv[1]) if len(sys.argv)>1 else ROOT/'docs/reports/FaultNet-final-report-revised.docx'
doc.save(path)
print(path)
print('Tables',table_num,'Figures',figure_num,'Headings',len(headings))
